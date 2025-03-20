import json
import  logging
from PyQt6 import QtWidgets, QtCore, QtGui
from PyQt6.QtCore import Qt, pyqtSignal
import openpyxl

log = logging.getLogger(__name__)

class DraggableTableWidget(QtWidgets.QTableWidget):
    rowsMoved = pyqtSignal()  # Add a signal

    def __init__(self, *args, config_file=None, **kwargs):
        super().__init__(*args, **kwargs)
        self._config_file = config_file
        self._do_not_save = True
        self.verticalHeader().setSectionsMovable(True)  # Enable row moving
        self.horizontalHeader().setSectionsMovable(True)  # Enable column moving
        self.verticalHeader().sectionMoved.connect(self._row_moved)
        self.horizontalHeader().sectionMoved.connect(self._col_moved)
        self.horizontalHeader().sectionResized.connect(self._col_resized)

        self._resize_timer = QtCore.QTimer(self)
        self._resize_timer.setSingleShot(True)
        self._resize_timer.timeout.connect(self.save_configuration)  # Connect directly to save_configuration

    def _row_moved(self, logicalIndex, oldVisualIndex, newVisualIndex):
        log.debug(f'row moved')
        self.rowsMoved.emit()

    def _col_moved(self, logicalIndex, oldVisualIndex, newVisualIndex):
        self.save_configuration()

    def _col_resized(self, logicalIndex=None, oldSize=None, newSize=None):
        if self._resize_timer.isActive():  # Check if the timer is running
            self._resize_timer.stop()  # If it is, stop it
        self._resize_timer.start(1000)  # Restart with the delay

    def setColumnHidden(self, column, hide):
        super().setColumnHidden(column, hide)
        if not hide and self._config:
            try:
                Width = self._config['column_widths'][column]
                Width = 100 if Width == 0 else Width
                super().setColumnWidth(column, Width)
            except:
                pass
        self.save_configuration()

    def save_configuration(self, filename=None):
        if self._do_not_save:
            return False
        fname = filename or self._config_file
        if not fname:
            log.error("Filename for storing table configuration not set")
            return False
        config = {}
        config["column_widths"] = [self.columnWidth(i) for i in range(self.columnCount())]
        config["column_visibility"] = [not self.isColumnHidden(i) for i in range(self.columnCount())]
        config["column_order"] = [self.horizontalHeader().visualIndex(i) for i in
                                  range(self.columnCount())]  # Get visual order

        try:
            with open(fname, 'w') as f:
                json.dump(config, f, indent=4)
                log.info(f"Table configuration saved to {fname}")
            self._config = config
            return True  # Or no return
        except Exception as e:  # Handle file writing errors
            log.error(f"Error saving configuration: {e}")  # Or logging or messagebox
            return False  # Or raise the error or return None

    def load_configuration(self, filename=None):
        fname = filename or self._config_file
        self._do_not_save = True
        if not fname:
            log.warning("Filename for loading table configuration not set")
            return False
        try:
            with open(fname, 'r') as f:
                config = json.load(f)
                log.debug(f"Table configuration loaded from {fname}")
        except FileNotFoundError:  # Handle file not found (maybe first time running).
            log.warning(f"Config file '{fname}' not found. Using default settings.")
            self._do_not_save = False
            return False
        except json.JSONDecodeError as e:  # Handle any json errors.
            log.error(f"Error loading configuration from '{fname}': Invalid JSON: {e}")
            self._do_not_save = False
            return False

        if config:
            # Restore column widths
            for i, width in enumerate(config.get("column_widths", [])):
                if i < self.columnCount():  # Ensure we don't go out of bounds
                    self.setColumnWidth(i, width)

            # Restore column visibility
            for i, visible in enumerate(config.get("column_visibility", [])):
                if i < self.columnCount():
                    if visible:
                        self.showColumn(i)
                    else:
                        self.hideColumn(i)

            # Restore column order (important to do this LAST)
            visual_order = config.get("column_order", [])
            if visual_order and len(visual_order) == self.columnCount():
                for i, visual_index in enumerate(visual_order):
                    self.horizontalHeader().moveSection(self.horizontalHeader().visualIndex(i), visual_index)

            self._config = config
        self._do_not_save = False

    def export_to_excel(self, filename):
        try:
            workbook = openpyxl.Workbook()
            sheet = workbook.active

            # Copy headers (both horizontal and vertical)
            for col in range(self.columnCount()):
                header_item = self.horizontalHeaderItem(col)
                if header_item:
                    visual_column_num = self.horizontalHeader().visualIndex(col)
                    sheet.cell(row=1, column=visual_column_num+2, value=header_item.text()) # Offset by 1 for vertical header

            for row in range(self.rowCount()):
                vertical_header_item = self.verticalHeaderItem(row)
                visual_row_num = self.verticalHeader().visualIndex(row)
                if vertical_header_item:
                     sheet.cell(row=visual_row_num + 2, column=1, value=vertical_header_item.text()) # Add vertical headers

                # Copy data (offset columns by 1 to accommodate vertical header)
                for col in range(self.columnCount()):
                    item = self.item(row, col)
                    visual_column_num = self.horizontalHeader().visualIndex(col)
                    if item:
                        sheet.cell(row=visual_row_num + 2, column=visual_column_num + 2, value=item.text())

            workbook.save(filename)
            log.info(f"Table exported to {filename}")
            return True

        except Exception as e:
            log.error(f"Error exporting to Excel: {e}")
            return False
